"""
Automated unit and integration tests for Zero AI Knowledge Base & RAG subsystem.
Tests text extraction across document types (PDF, DOCX, PPTX, XLSX, CSV, Code, Text),
audio/video rejection, empty/malformed files, chunking, FTS5 BM25 search,
incremental reindexing, and deletion handling.
"""

import tempfile
import os
from pathlib import Path
import pytest
from app.knowledge.parsers import DocumentParserFactory
from app.knowledge.chunker import TextChunker
from app.knowledge.store import KnowledgeStore, SearchResult
from app.knowledge.service import KnowledgeService


@pytest.fixture
def temp_kb_dir():
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def temp_store():
    with tempfile.TemporaryDirectory() as tmpdir:
        db_path = Path(tmpdir) / "test_knowledge.db"
        store = KnowledgeStore(db_path=str(db_path))
        yield store


def test_supported_extensions_and_media_rejection():
    """Verify supported extensions and rejection of audio/video files."""
    supported = DocumentParserFactory.get_supported_extensions()
    assert ".pdf" in supported
    assert ".docx" in supported
    assert ".pptx" in supported
    assert ".xlsx" in supported
    assert ".py" in supported
    assert ".txt" in supported
    assert ".md" in supported

    # Verify audio/video files are explicitly rejected
    for media_ext in [".mp3", ".wav", ".m4a", ".mp4", ".avi", ".mov", ".mkv"]:
        assert not DocumentParserFactory.can_parse(Path(f"test{media_ext}"))


def test_text_and_code_extraction(temp_kb_dir):
    """Test extracting plain text and programming source code."""
    txt_file = temp_kb_dir / "notes.txt"
    txt_file.write_text("Zero AI is an advanced personal assistant.\nIt uses local retrieval.", encoding="utf-8")
    parsed_txt = DocumentParserFactory.parse(txt_file)
    assert parsed_txt.success
    assert "Zero AI is an advanced personal assistant" in parsed_txt.text

    py_file = temp_kb_dir / "script.py"
    py_file.write_text("def assist():\n    print('Standing by.')\n", encoding="utf-8")
    parsed_py = DocumentParserFactory.parse(py_file)
    assert parsed_py.success
    assert "def assist():" in parsed_py.text


def test_csv_and_json_extraction(temp_kb_dir):
    """Test extracting CSV tabular data and JSON structures."""
    csv_file = temp_kb_dir / "data.csv"
    csv_file.write_text("Name,Role,Status\nZero,Assistant,Active\nJarvis,Core,Standby\n", encoding="utf-8")
    parsed_csv = DocumentParserFactory.parse(csv_file)
    assert parsed_csv.success
    assert "Zero" in parsed_csv.text
    assert "Assistant" in parsed_csv.text

    json_file = temp_kb_dir / "config.json"
    json_file.write_text('{"app": "Zero", "mode": "RAG", "port": 8765}', encoding="utf-8")
    parsed_json = DocumentParserFactory.parse(json_file)
    assert parsed_json.success
    assert "RAG" in parsed_json.text


def test_docx_extraction(temp_kb_dir):
    """Test DOCX word document text extraction."""
    from docx import Document
    doc = Document()
    doc.add_heading("Zero AI System Manual", level=1)
    doc.add_paragraph("This paragraph describes the knowledge base architecture.")
    docx_path = temp_kb_dir / "manual.docx"
    doc.save(str(docx_path))

    parsed = DocumentParserFactory.parse(docx_path)
    assert parsed.success
    assert "Zero AI System Manual" in parsed.text
    assert "knowledge base architecture" in parsed.text


def test_pptx_extraction(temp_kb_dir):
    """Test PPTX presentation text extraction."""
    from pptx import Presentation
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Project Zero Overview"
    subtitle.text = "Futuristic JARVIS Desktop Assistant"
    pptx_path = temp_kb_dir / "presentation.pptx"
    prs.save(str(pptx_path))

    parsed = DocumentParserFactory.parse(pptx_path)
    assert parsed.success
    assert "Project Zero Overview" in parsed.text
    assert "Futuristic JARVIS Desktop Assistant" in parsed.text


def test_xlsx_extraction(temp_kb_dir):
    """Test XLSX spreadsheet text extraction."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Telemetry"
    ws.append(["Metric", "Value", "Unit"])
    ws.append(["Latency", "4.2", "ms"])
    xlsx_path = temp_kb_dir / "metrics.xlsx"
    wb.save(str(xlsx_path))

    parsed = DocumentParserFactory.parse(xlsx_path)
    assert parsed.success
    assert "Latency" in parsed.text
    assert "Telemetry" in parsed.sections[0].label


def test_pdf_extraction(temp_kb_dir):
    """Test PDF text extraction."""
    from pypdf import PdfWriter
    writer = PdfWriter()
    # Create empty page or text page using writer
    writer.add_blank_page(width=72, height=72)
    pdf_path = temp_kb_dir / "document.pdf"
    with pdf_path.open("wb") as f:
        writer.write(f)

    parsed = DocumentParserFactory.parse(pdf_path)
    assert parsed.success


def test_empty_and_malformed_files(temp_kb_dir):
    """Test empty and corrupted files do not crash the parser."""
    empty_file = temp_kb_dir / "empty.txt"
    empty_file.write_text("", encoding="utf-8")
    parsed_empty = DocumentParserFactory.parse(empty_file)
    assert parsed_empty.text == ""
    assert parsed_empty.error == "File is empty (0 bytes)."

    corrupted_pdf = temp_kb_dir / "corrupted.pdf"
    corrupted_pdf.write_bytes(b"NOT A REAL PDF FILE")
    parsed_corrupt = DocumentParserFactory.parse(corrupted_pdf)
    assert not parsed_corrupt.success
    assert parsed_corrupt.error is not None


def test_chunker_overlapping_and_metadata():
    """Verify text chunker splits content with overlap and metadata."""
    chunker = TextChunker(chunk_size=100, chunk_overlap=20)
    text = "Zero AI assistant provides seamless voice interactions, local knowledge retrieval, and PDF document generation. " * 5
    chunks = chunker.chunk_text(text, source_file="overview.md", section_label="Overview")

    assert len(chunks) > 1
    for chunk in chunks:
        assert chunk.source_file == "overview.md"
        assert chunk.section_label == "Overview"
        assert len(chunk.content) > 0
        assert chunk.content_hash is not None


def test_fts5_indexing_and_search(temp_store):
    """Test full-text search and BM25 ranking in SQLite."""
    chunker = TextChunker(chunk_size=200, chunk_overlap=30)
    c1 = chunker.chunk_text(
        "Zero AI uses Playwright to compile executive HTML documents into verified PDF reports.",
        source_file="pdf_spec.md",
        section_label="Section 1",
    )
    c2 = chunker.chunk_text(
        "The voice engine coordinates SpeechRecognition with background pyttsx3 speech synthesis.",
        source_file="voice_spec.md",
        section_label="Section 2",
    )

    temp_store.index_document(
        file_id="pdf_spec.md",
        file_path="/Knowledge/pdf_spec.md",
        file_hash="hash1",
        file_size=1024,
        mtime=1000.0,
        chunks=c1,
    )
    temp_store.index_document(
        file_id="voice_spec.md",
        file_path="/Knowledge/voice_spec.md",
        file_hash="hash2",
        file_size=2048,
        mtime=2000.0,
        chunks=c2,
    )

    # Search for Playwright PDF
    results = temp_store.search("Playwright PDF", limit=5)
    assert len(results) > 0
    assert results[0].source_file == "pdf_spec.md"
    assert "Playwright" in results[0].content

    # Search for speech synthesis
    voice_results = temp_store.search("speech synthesis", limit=5)
    assert len(voice_results) > 0
    assert voice_results[0].source_file == "voice_spec.md"


def test_incremental_reindexing_and_deletion(temp_kb_dir, temp_store):
    """Test that unchanged files are skipped, updated files are reindexed, and deleted files pruned."""
    service = KnowledgeService(knowledge_dir=temp_kb_dir, store=temp_store)

    doc1 = temp_kb_dir / "doc1.txt"
    doc1.write_text("First version of document one content for testing.", encoding="utf-8")

    # Initial indexing run
    r1 = service.sync_and_index(force_all=False)
    assert r1.indexed_files == 1
    assert r1.skipped_files == 0

    # Second run without changes -> should skip (incremental)
    r2 = service.sync_and_index(force_all=False)
    assert r2.indexed_files == 0
    assert r2.skipped_files == 1

    # Modify doc1 -> should update
    doc1.write_text("Second modified version of document one with new terms.", encoding="utf-8")
    r3 = service.sync_and_index(force_all=False)
    assert r3.updated_files == 1

    # Delete doc1 from disk -> should prune from database
    doc1.unlink()
    r4 = service.sync_and_index(force_all=False)
    assert r4.deleted_files == 1
    assert len(service.list_files()) == 0


def test_rag_context_formatting_and_security(temp_store):
    """Verify RAG prompt formatting encloses document excerpts as untrusted reference material."""
    chunker = TextChunker()
    chunks = chunker.chunk_text(
        "Important Project Data: Project Codename is AEGIS-7.",
        source_file="project.txt",
        section_label="Confidential",
    )
    temp_store.index_document(
        file_id="project.txt",
        file_path="/Knowledge/project.txt",
        file_hash="hash_p",
        file_size=512,
        mtime=100.0,
        chunks=chunks,
    )

    service = KnowledgeService(store=temp_store)
    prompt_context, search_results = service.format_rag_context("AEGIS-7", limit=3)

    assert len(search_results) == 1
    assert "[KNOWLEDGE BASE RETRIEVAL ACTIVE]" in prompt_context
    assert "[Doc: project.txt" in prompt_context or "Source: project.txt" in prompt_context
    assert "AEGIS-7" in prompt_context
    assert "untrusted reference material" in prompt_context or "informational data" in prompt_context
